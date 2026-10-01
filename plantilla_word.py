from __future__ import annotations

import re
import shutil
import time
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from pdf_extractor import DatosAutoAdmite

BASE_DIR = Path(__file__).resolve().parent
PLANTILLA_DEFAULT = BASE_DIR / "Plantilla Autonomia Universitaria.docx"
VINCULADOS_FIJO = "MINISTERIO DE EDUCACIÓN NACIONAL"
FUENTE_DOCUMENTO = "Verdana"

_ORDEN_PRETENSIONES = (
    "PRIMERA",
    "SEGUNDA",
    "TERCERA",
    "CUARTA",
    "QUINTA",
    "SEXTA",
    "SÉPTIMA",
)


def _estilo_run(run, *, negrita: bool = False, cursiva: bool = False) -> None:
    run.font.name = FUENTE_DOCUMENTO
    run.font.bold = negrita
    run.font.italic = cursiva
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), FUENTE_DOCUMENTO)
    r_fonts.set(qn("w:hAnsi"), FUENTE_DOCUMENTO)
    r_fonts.set(qn("w:cs"), FUENTE_DOCUMENTO)
    if negrita:
        r_pr.get_or_add_b()
    else:
        bold = r_pr.find(qn("w:b"))
        if bold is not None:
            r_pr.remove(bold)
    if cursiva:
        r_pr.get_or_add_i()
    else:
        italic = r_pr.find(qn("w:i"))
        if italic is not None:
            r_pr.remove(italic)
    run.font.highlight_color = None
    resaltado = r_pr.find(qn("w:highlight"))
    if resaltado is not None:
        r_pr.remove(resaltado)


def _limpiar_parrafo(parrafo: Paragraph) -> None:
    element = parrafo._element
    for child in list(element):
        if child.tag == qn("w:r"):
            element.remove(child)


def _asignar_campo_cabecera(
    parrafo: Paragraph,
    etiqueta: str,
    valor: str,
    separador: str,
) -> None:
    _limpiar_parrafo(parrafo)
    run_etiqueta = parrafo.add_run(etiqueta)
    _estilo_run(run_etiqueta, negrita=True)
    run_valor = parrafo.add_run(f"{separador}{valor}")
    _estilo_run(run_valor, negrita=False)


def _aplicar_verdana_documento(doc: Document) -> None:
    for indice, parrafo in enumerate(doc.paragraphs):
        ordinal = _ordinal_desde_parrafo(parrafo.text)
        es_pretension = ordinal in _ORDEN_PRETENSIONES
        for posicion, run in enumerate(parrafo.runs):
            if not run.text:
                continue
            _estilo_run(
                run,
                negrita=indice < 5 and posicion == 0,
                cursiva=es_pretension,
            )
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                for parrafo in celda.paragraphs:
                    for run in parrafo.runs:
                        if run.text:
                            _estilo_run(run, negrita=False, cursiva=False)


def _asignar_texto_conservando_estilo(
    parrafo: Paragraph,
    texto: str,
    *,
    cursiva: bool = False,
) -> None:
    _limpiar_parrafo(parrafo)
    run = parrafo.add_run(texto)
    _estilo_run(run, negrita=False, cursiva=cursiva)


def _reemplazar_accionado_radicado_externo(
    parrafo: Paragraph,
    accionado: str,
) -> bool:
    """Sustituye el nombre de ejemplo tras «radicado ante la …»."""
    if not accionado.strip():
        return False
    texto = parrafo.text
    coincidencia = re.search(
        r"(?is)(.+?\bhaberlo radicado ante la\s+)(.+?)(\s*,\s*entidad diferente)",
        texto,
    )
    if not coincidencia:
        return False

    prefijo = coincidencia.group(1)
    sufijo = coincidencia.group(3) + texto[coincidencia.end() :]

    _limpiar_parrafo(parrafo)
    run_prefijo = parrafo.add_run(prefijo)
    _estilo_run(run_prefijo, negrita=False)
    run_accionado = parrafo.add_run(accionado.strip())
    _estilo_run(run_accionado, negrita=False)
    run_sufijo = parrafo.add_run(sufijo)
    _estilo_run(run_sufijo, negrita=False)
    return True


def _reemplazar_fecha_en_parrafo(parrafo: Paragraph, fecha: str) -> None:
    if not fecha:
        return
    texto = parrafo.text
    nuevo = re.sub(
        r"(\d{1,2}\s+de\s+[a-záéíóúñA-ZÁÉÍÓÚÑ]+\s+de\s+\d{4})",
        fecha,
        texto,
        count=1,
        flags=re.IGNORECASE,
    )
    if nuevo != texto:
        _asignar_texto_conservando_estilo(parrafo, nuevo)


def _prefijo_pretension(texto_parrafo: str) -> str:
    coincidencia = re.match(r"^([^\w]*\w+\.\s*)", texto_parrafo)
    if coincidencia:
        return coincidencia.group(1)
    return ""


def _ordinal_desde_parrafo(texto_parrafo: str) -> str | None:
    coincidencia = re.search(
        r"\b(PRIMERA|SEGUNDA|TERCERA|CUARTA|QUINTA|SEXTA|S[ÉE]PTIMA)\.",
        texto_parrafo,
        re.IGNORECASE,
    )
    if not coincidencia:
        return None
    ordinal = coincidencia.group(1).upper().replace("SEPTIMA", "SÉPTIMA")
    return ordinal


def _archivo_en_uso(exc: OSError) -> bool:
    return exc.errno in (13, 32) or getattr(exc, "winerror", None) in (5, 32)


def _publicar_docx_generado(temporal: Path, destino: Path) -> Path:
    """Mueve el .docx temporal al destino; reintenta si Word lo tiene abierto."""
    for intento in range(12):
        try:
            temporal.replace(destino)
            return destino
        except OSError as exc:
            if not _archivo_en_uso(exc):
                raise
            time.sleep(0.6 * (intento + 1))

    marca = time.strftime("%Y%m%d_%H%M%S")
    alterno = destino.with_name(f"{destino.stem}_{marca}{destino.suffix}")
    temporal.replace(alterno)
    return alterno


def llenar_plantilla(
    datos: DatosAutoAdmite,
    ruta_salida: Path,
    *,
    plantilla: Path | None = None,
) -> Path:
    origen = plantilla or PLANTILLA_DEFAULT
    if not origen.is_file():
        raise FileNotFoundError(f"No se encontró la plantilla: {origen}")

    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    temporal = ruta_salida.with_name(f"{ruta_salida.stem}.__tmp__.docx")
    shutil.copy2(origen, temporal)
    doc = Document(str(temporal))

    if doc.paragraphs:
        _asignar_campo_cabecera(
            doc.paragraphs[0], "RADICADO:", datos.radicado, "\t\t"
        )
    if len(doc.paragraphs) > 1 and datos.tipo_proceso:
        _asignar_campo_cabecera(
            doc.paragraphs[1],
            "TIPO DE PROCESO: ",
            f"{datos.tipo_proceso} ",
            "\t",
        )
    if len(doc.paragraphs) > 2:
        _asignar_campo_cabecera(
            doc.paragraphs[2],
            "ACCIONANTE:",
            datos.accionante,
            "               ",
        )
    if len(doc.paragraphs) > 3:
        _asignar_campo_cabecera(
            doc.paragraphs[3], "ACCIONADO:", datos.accionado, "\t"
        )
    if len(doc.paragraphs) > 4:
        _asignar_campo_cabecera(
            doc.paragraphs[4], "VINCULADOS:", f"{VINCULADOS_FIJO} ", "\t"
        )

    for indice in (10, 14, 18):
        if len(doc.paragraphs) > indice:
            _reemplazar_fecha_en_parrafo(doc.paragraphs[indice], datos.fecha_auto_admisorio)

    for parrafo in doc.paragraphs:
        ordinal = _ordinal_desde_parrafo(parrafo.text)
        if not ordinal or ordinal not in _ORDEN_PRETENSIONES:
            continue
        cuerpo = datos.pretensiones.get(ordinal)
        if not cuerpo:
            _limpiar_parrafo(parrafo)
            continue
        prefijo = _prefijo_pretension(parrafo.text) or f"{ordinal}. "
        _asignar_texto_conservando_estilo(
            parrafo, f"{prefijo}{cuerpo}", cursiva=True
        )

    if datos.accionado:
        for parrafo in doc.paragraphs:
            if _reemplazar_accionado_radicado_externo(parrafo, datos.accionado):
                break

    _aplicar_verdana_documento(doc)
    doc.save(str(temporal))
    publicado = _publicar_docx_generado(temporal, ruta_salida)
    return publicado
